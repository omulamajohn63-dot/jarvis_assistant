"""Reading a document's text out of whatever format it is saved in.

Each reader is imported inside its own function, so a missing package stops one
format from being read rather than stopping the assistant from starting.
"""

from __future__ import annotations

import os
import re

from jarvis2.logging_utils import get_logger

_log = get_logger("files.documents")

# How much of a document is read. A whole document can run to hundreds of pages,
# which would take longer to say than to read.
READ_CHAR_LIMIT = 1200

# How many pages of a PDF to pull text from, since text is read in order.
READ_PAGES = 3


def plain_text(path: str) -> str:
    """Reads a text file, replacing any characters it cannot decode."""
    with open(path, encoding="utf-8", errors="replace") as handle:
        return handle.read()


def rtf_text(path: str) -> str:
    """Reads an RTF file, dropping the formatting codes it is full of.

    Spoken as it stands, the codes come out as noise, so they are stripped
    before anything is said.
    """
    text = re.sub(r"\\[a-z]+-?\d* ?", " ", plain_text(path))
    return re.sub(r"[{}]", " ", text)


def pdf_text(path: str) -> str:
    """Reads the opening pages of a PDF."""
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n".join(
        page.extract_text() or "" for page in reader.pages[:READ_PAGES]
    )


def docx_text(path: str) -> str:
    """Reads a Word document's paragraphs and table cells."""
    import docx

    document = docx.Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def pptx_text(path: str) -> str:
    """Reads the text on each slide of a PowerPoint deck."""
    from pptx import Presentation

    parts = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def xlsx_text(path: str) -> str:
    """Reads a spreadsheet's cells, sheet by sheet."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for sheet in workbook.worksheets:
            parts.append(sheet.title)
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if cells:
                    parts.append(" ".join(cells))
        return "\n".join(parts)
    finally:
        workbook.close()


# The formats that can be read. Anything not here has no reader, and Jarvis
# says so and opens the file instead of failing.
READERS = {
    ".txt": plain_text,
    ".md": plain_text,
    ".csv": plain_text,
    ".log": plain_text,
    ".rtf": rtf_text,
    ".pdf": pdf_text,
    ".docx": docx_text,
    ".pptx": pptx_text,
    ".xlsx": xlsx_text,
    ".xlsm": xlsx_text,
}

# Recognised, but with no reader: the pre-2007 formats.
UNREADABLE = {".doc", ".xls", ".ppt"}


def can_read(path: str) -> bool:
    return os.path.splitext(path)[1].lower() in READERS


def extension_name(path: str) -> str:
    """The format's name as a person would say it, or an empty string."""
    extension = os.path.splitext(path)[1].lower()
    return extension.lstrip(".").upper() if extension else ""


def read(path: str) -> str:
    """A document's text, or an empty string when it cannot be read.

    One unreadable file does not take the assistant down: the reason is logged
    and the caller is told there is nothing to read.
    """
    reader = READERS.get(os.path.splitext(path)[1].lower())
    if reader is None:
        return ""

    try:
        return reader(path)
    except Exception as error:
        _log.warning("could not read %s: %s", os.path.basename(path), error)
        return ""


def prepare(path: str, limit: int = READ_CHAR_LIMIT) -> tuple:
    """A document's opening, ready to be spoken.

    Returns (text, truncated). Whitespace is squashed first, because documents
    are full of line breaks and runs of spaces which read aloud as pauses.
    """
    collapsed = " ".join(read(path).split())
    if not collapsed:
        return "", False
    truncated = len(collapsed) > limit
    return collapsed[:limit].strip(), truncated
